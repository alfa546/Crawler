import json
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from flask import send_file, request, jsonify
from . import crawler_bp

@crawler_bp.route('/export-crawl-xlsx', methods=['POST'])
def export_crawl_xlsx():
    """Export crawl results as a styled .xlsx workbook.

    Sheet 1 = All Pages (every URL, key SEO fields, color-coded status/speed).
    Sheet 2 = Issues Summary (issue text → page count → sample URLs).
    Sheet 3+ = whatever the client supplied in `extra_sheets` (built by
    `_buildExportForCategory` on the frontend so each tab gets the shape it
    actually shows).
    """
    import io as _io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    data = request.json or {}
    results = data.get('results', [])
    domain = data.get('domain', 'site')

    if not results:
        return jsonify({'error': 'No results to export'}), 400

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'All Pages'

    headers = ['URL', 'Status', 'Title', 'Title Len', 'Meta Description', 'Meta Len',
               'H1', 'Words', 'Canonical', 'Indexable', 'Int Links', 'Ext Links',
               'Images', 'Alt Missing', 'Schema', 'Speed (s)', 'Depth', 'Issues']

    header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='6B5CE7', end_color='6B5CE7', fill_type='solid')
    cell_font = Font(name='Calibri', size=10)
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'), right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'), bottom=Side(style='thin', color='D0D0D0'),
    )
    green_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
    amber_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    red_fill   = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for ri, r in enumerate(results, 2):
        vals = [
            r.get('url', ''), r.get('status_code', 0), r.get('title', ''), r.get('title_len', 0),
            r.get('meta_description', ''), r.get('meta_len', 0), r.get('h1', ''),
            r.get('word_count', 0), r.get('canonical', ''),
            'Yes' if r.get('indexable', True) else 'No',
            r.get('internal_links', 0), r.get('external_links', 0),
            r.get('images_total', 0), r.get('images_no_alt', 0),
            ', '.join(r.get('schema_types', [])[:3]),
            r.get('response_time', 0), r.get('depth', 0),
            '; '.join(r.get('issues', []))
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font = cell_font
            cell.border = thin_border
        status = r.get('status_code', 0)
        sc = ws.cell(row=ri, column=2)
        if 200 <= status < 300: sc.fill = green_fill
        elif 300 <= status < 400: sc.fill = amber_fill
        elif status >= 400: sc.fill = red_fill
        speed = r.get('response_time', 0)
        sp = ws.cell(row=ri, column=16)
        if speed <= 1: sp.fill = green_fill
        elif speed <= 3: sp.fill = amber_fill
        else: sp.fill = red_fill

    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, min(len(str(cell.value)), 60))
        ws.column_dimensions[letter].width = max_len + 3
    ws.freeze_panes = 'A2'

    # Sheet 2: Issues Summary
    ws2 = wb.create_sheet('Issues Summary')
    for ci, h in enumerate(['Issue', 'Count', 'Pages'], 1):
        cell = ws2.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
    issue_map = {}
    for r in results:
        for issue in r.get('issues', []):
            base = issue.split('(')[0].strip()
            issue_map.setdefault(base, []).append(r.get('url', ''))
    for ri, (issue, urls) in enumerate(sorted(issue_map.items(), key=lambda x: -len(x[1])), 2):
        ws2.cell(row=ri, column=1, value=issue).font = cell_font
        ws2.cell(row=ri, column=2, value=len(urls)).font = Font(name='Calibri', size=10, bold=True)
        ws2.cell(row=ri, column=3, value='; '.join(urls[:10])).font = cell_font
    ws2.column_dimensions['A'].width = 30
    ws2.column_dimensions['B'].width = 8
    ws2.column_dimensions['C'].width = 80

    # Sheets 3+ — whatever the client built. {name, header, rows}.
    extra = data.get('extra_sheets') or []
    used_names = {ws.title, ws2.title}
    def _safe_sheet_name(n):
        s = _re.sub(r'[:\\/\?\*\[\]]', '-', str(n or 'Sheet'))[:31].strip() or 'Sheet'
        if s in used_names:
            base = s[:28]
            i = 2
            while f'{base} {i}' in used_names and i < 99:
                i += 1
            s = f'{base} {i}'
        used_names.add(s)
        return s
    for sheet in extra:
        name = _safe_sheet_name(sheet.get('name') or 'Sheet')
        header = sheet.get('header') or []
        rows = sheet.get('rows') or []
        if not header and not rows:
            continue
        ws_e = wb.create_sheet(name)
        for ci, h in enumerate(header, 1):
            cell = ws_e.cell(row=1, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        for ri, r in enumerate(rows, 2):
            for ci, v in enumerate(r, 1):
                if isinstance(v, (list, tuple)):
                    v = ', '.join(str(x) for x in v)
                elif isinstance(v, dict):
                    v = json.dumps(v, ensure_ascii=False)
                cell = ws_e.cell(row=ri, column=ci, value=v)
                cell.font = cell_font
                cell.border = thin_border
        for col in ws_e.columns:
            max_len = 0
            letter = col[0].column_letter
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 80))
            ws_e.column_dimensions[letter].width = max(8, max_len + 2)
        ws_e.freeze_panes = 'A2'

    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_domain = _re.sub(r'[^a-zA-Z0-9.-]', '', (domain or '').lstrip('.').replace('www.', '', 1))
    from datetime import datetime as _dt
    _ts = _dt.now().strftime('%Y-%m-%d-%H%M')
    _name = '-'.join([p for p in ['crawl', safe_domain, _ts] if p]) + '.xlsx'
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=_name,
    )


@crawler_bp.route('/export-crawl-sitemap', methods=['POST'])
def export_crawl_sitemap():
    """Build a sitemaps.org 0.9 XML sitemap from the crawl results.

    Filter: include only URLs that a search engine would actually want to
    index - HTTP 200, indexable (no noindex meta or X-Robots-Tag), self- or
    missing-canonical (skip pages that canonical to a different URL since
    those aren't the canonical version), and not redirected. >50K URLs
    splits into a sitemap-index plus N child sitemaps to stay within the
    sitemaps.org 50,000-URL / 50MB-uncompressed limits.
    """
    from xml.sax.saxutils import escape as _xml_escape
    from datetime import datetime as _dt
    from email.utils import parsedate_to_datetime as _parsedate
    import io as _io
    import zipfile as _zipfile

    data = request.json or {}
    results = data.get('results', [])
    domain = (data.get('domain') or '').strip()

    if not results:
        return json.dumps({'error': 'No results to export'}), 400

    def _is_indexable_200(r):
        if r.get('status_code') != 200:
            return False
        if r.get('error'):
            return False
        if r.get('redirect_url'):
            return False
        if r.get('indexable') is False:
            return False
        if r.get('canonical_kind') == 'canonicalised':
            return False
        ctype = (r.get('content_type') or '').lower()
        if ctype and 'html' not in ctype and 'xml' not in ctype:
            return False
        if r.get('is_pagination'):
            return False
        url = (r.get('url') or '').strip()
        if not url or not (url.startswith('http://') or url.startswith('https://')):
            return False
        return True

    eligible = [r for r in results if _is_indexable_200(r)]

    seen = set()
    urls = []
    for r in eligible:
        u = r['url'].strip()
        if u in seen:
            continue
        seen.add(u)
        # Format Last-Modified as W3C date (YYYY-MM-DD). Header arrives as
        # RFC 1123; fall back silently if missing or malformed.
        lastmod = ''
        lm_raw = (r.get('last_modified') or '').strip()
        if lm_raw:
            try:
                lastmod = _parsedate(lm_raw).strftime('%Y-%m-%d')
            except Exception:
                lastmod = ''
        urls.append({'loc': u, 'lastmod': lastmod})

    SITEMAP_NS = 'http://www.sitemaps.org/schemas/sitemap/0.9'
    URLS_PER_SITEMAP = 50000

    def _build_urlset(chunk):
        lines = ['<?xml version="1.0" encoding="UTF-8"?>',
                 f'<urlset xmlns="{SITEMAP_NS}">']
        for u in chunk:
            lines.append('  <url>')
            lines.append(f'    <loc>{_xml_escape(u["loc"])}</loc>')
            if u['lastmod']:
                lines.append(f'    <lastmod>{u["lastmod"]}</lastmod>')
            lines.append('  </url>')
        lines.append('</urlset>')
        lines.append('')
        return '\n'.join(lines)

    safe_domain = _re.sub(r'[^a-zA-Z0-9.-]', '', (domain or '').lstrip('.').replace('www.', '', 1))
    _ts = _dt.now().strftime('%Y-%m-%d-%H%M')

    if len(urls) <= URLS_PER_SITEMAP:
        xml = _build_urlset(urls)
        _name = '-'.join([p for p in ['sitemap', safe_domain, _ts] if p]) + '.xml'
        return send_file(_io.BytesIO(xml.encode('utf-8')),
            mimetype='application/xml',
            as_attachment=True,
            download_name=_name)

    zip_buf = _io.BytesIO()
    today = _dt.now().strftime('%Y-%m-%d')
    with _zipfile.ZipFile(zip_buf, 'w', _zipfile.ZIP_DEFLATED) as zf:
        index_lines = ['<?xml version="1.0" encoding="UTF-8"?>',
                       f'<sitemapindex xmlns="{SITEMAP_NS}">']
        for i in range(0, len(urls), URLS_PER_SITEMAP):
            chunk = urls[i:i + URLS_PER_SITEMAP]
            child_name = f'sitemap-{i // URLS_PER_SITEMAP + 1}.xml'
            zf.writestr(child_name, _build_urlset(chunk))
            host = domain.rstrip('/') if domain else ''
            index_lines.append('  <sitemap>')
            index_lines.append(f'    <loc>{_xml_escape(host + "/" + child_name)}</loc>')
            index_lines.append(f'    <lastmod>{today}</lastmod>')
            index_lines.append('  </sitemap>')
        index_lines.append('</sitemapindex>')
        index_lines.append('')
        zf.writestr('sitemap.xml', '\n'.join(index_lines))
    zip_buf.seek(0)
    _name = '-'.join([p for p in ['sitemap', safe_domain, _ts] if p]) + '.zip'
    return send_file(zip_buf,
        mimetype='application/zip',
        as_attachment=True,
        download_name=_name)


